from __future__ import annotations

import ast
import contextlib
import io
import os
import re
import traceback
from pathlib import Path
from typing import Any, Dict, Optional, Tuple

import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Mm, Pt, RGBColor


class PythonDocxCodeGenerator:
    """
    Generates and executes concise, robust python-docx code for file manipulation tasks.
    Supports creating new documents, precision editing with Paragraph IDs ([P0], [P1]...),
    advanced image wrapping/layering, and custom table/page borders.
    """

    CONTEXT = r"""
# Python-DOCX Reference & Best Practices

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# 1. Opening Documents:
# A) Brand NEW document:
# filename = 'new_doc.docx'
# doc = Document()

# B) EDITING an EXISTING document:
# filename = 'existing_doc.docx'
# doc = Document(filename)

# 2. Table Borders (Black, Visible, or Custom Colors):
# - Setting standard visible black grid borders on table [T0]:
#   table = doc.tables[0]
#   table.style = 'Table Grid'
#   # Or via explicit OpenXML tblBorders:
#   tblPr = table._tbl.tblPr
#   for old_b in tblPr.xpath('./w:tblBorders'): tblPr.remove(old_b)
#   tblBorders = parse_xml(f'''
#   <w:tblBorders {nsdecls("w")}>
#       <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>
#       <w:left w:val="single" w:sz="8" w:space="0" w:color="000000"/>
#       <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>
#       <w:right w:val="single" w:sz="8" w:space="0" w:color="000000"/>
#       <w:insideH w:val="single" w:sz="8" w:space="0" w:color="000000"/>
#       <w:insideV w:val="single" w:sz="8" w:space="0" w:color="000000"/>
#   </w:tblBorders>''')
#   tblPr.append(tblBorders)

# 3. Image Layering & Text Wrapping Controls:
# - INLINE Image (Default):
#   p_img = doc.add_paragraph()
#   p_img.add_run().add_picture(img_path, width=Inches(3.5))

# - FLOATING Image: Behind Text (Underlay / Watermark) or In Front of Text (Overlay):
#   def make_floating_picture(run, behind_text=False, wrap='none'):
#       inline = run._r.xpath('.//wp:inline')
#       if not inline: return
#       inline_el = inline[0]
#       extent = inline_el.xpath('.//wp:extent')[0]
#       graphic = inline_el.xpath('.//a:graphic')[0]
#       cx, cy = extent.get('cx'), extent.get('cy')
#       behind_val = '1' if behind_text else '0'
#       wrap_tag = '<wp:wrapSquare/>' if wrap == 'square' else '<wp:wrapNone/>'
#       anchor_xml = f'''
#       <wp:anchor {nsdecls("wp")} distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251658240" behindDoc="{behind_val}" locked="0" layoutInCell="1" allowOverlap="1">
#         <wp:simplePos x="0" y="0"/>
#         <wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>
#         <wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>
#         <wp:extent cx="{cx}" cy="{cy}"/>
#         <wp:effectExtent l="0" t="0" r="0" b="0"/>
#         {wrap_tag}
#         <wp:docPr id="1" name="Picture"/>
#       </wp:anchor>'''
#       anchor = parse_xml(anchor_xml)
#       anchor.append(graphic)
#       drawing = inline_el.getparent()
#       drawing.remove(inline_el)
#       drawing.append(anchor)

# 4. Page Borders:
# - Standard line borders or graphic/art borders ('stars', 'certificateBanner', 'vine'):
#   ns = nsdecls('w')
#   for section in doc.sections:
#       sectPr = section._sectPr
#       for pb in sectPr.xpath('./w:pgBorders'): sectPr.remove(pb)
#       sectPr.append(parse_xml(f'''
#       <w:pgBorders {ns} w:offsetFrom="text">
#           <w:top w:val="single" w:sz="24" w:space="24" w:color="000000"/>
#           <w:left w:val="single" w:sz="24" w:space="24" w:color="000000"/>
#           <w:bottom w:val="single" w:sz="24" w:space="24" w:color="000000"/>
#           <w:right w:val="single" w:sz="24" w:space="24" w:color="000000"/>
#       </w:pgBorders>'''))

# 5. Saving:
# doc.save(filename)
"""

    def _extract_clean_code(self, raw_content: str) -> str:
        """Extract clean Python code from model response, stripping thinking tags and code fences."""
        if not raw_content:
            return ""

        text = raw_content.strip()
        text = re.sub(r"<think>[\s\S]*?</think>", "", text, flags=re.IGNORECASE).strip()

        if "```" in text:
            match = re.search(r"```(?:python)?\s*([\s\S]*?)(?:```|$)", text, re.IGNORECASE)
            if match:
                text = match.group(1).strip()

        lines = text.splitlines()
        code_start_keywords = (
            "import ", "from ", "#", "doc =", "doc=", "def ", "class ",
            "table =", "table=", "p =", "p=", "run =", "run=", "filename =", "filename=", "p_img ="
        )
        start_idx = 0
        found_code = False
        for idx, line in enumerate(lines):
            l_str = line.strip()
            if any(l_str.startswith(kw) for kw in code_start_keywords):
                start_idx = idx
                found_code = True
                break

        if found_code and start_idx > 0:
            text = "\n".join(lines[start_idx:]).strip()

        return text

    def generate(
        self,
        task_description: str,
        nemotron_client: Any,
        model: str = "nvidia/nemotron-3-super-120b-a12b",
        error_context: Optional[str] = None,
    ) -> str:
        code, _ = self.generate_with_usage(task_description, nemotron_client, model, error_context)
        return code

    def generate_with_usage(
        self,
        task_description: str,
        nemotron_client: Any,
        model: str = "nvidia/nemotron-3-super-120b-a12b",
        error_context: Optional[str] = None,
    ) -> Tuple[str, int]:
        """
        Generate concise, executable python-docx code and track token usage.
        """
        system_prompt = f"""You are a python-docx expert.

Generate clean, accurate, production-ready Python code using the python-docx library.

{self.CONTEXT}

CRITICAL RULES:
1. Generate ONLY valid, executable Python code with all necessary imports.
2. TABLE BORDERS: To make table borders visible/black, set `table.style = 'Table Grid'` or append `<w:tblBorders>` XML to `table._tbl.tblPr`.
3. PARAGRAPH & TABLE IDENTIFIERS ([P0], [P1], [T0]...): Use exact indices `doc.paragraphs[idx]` or `doc.tables[idx]` from the provided Document Map.
4. PRESERVE EXISTING CONTENT: When editing, load with `doc = Document(filename)` and modify in-place.
5. Always save the document at the end with `doc.save(filename)`.
6. No conversational explanations, output ONLY executable Python code."""

        user_content = task_description
        if error_context:
            user_content += f"\n\nIMPORTANT: Fix this previous error:\n{error_context}"

        tokens_used = 0
        kwargs: Dict[str, Any] = {
            "model": model,
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_content},
            ],
            "temperature": 0.1,
            "max_tokens": 2048,
        }

        try:
            response = nemotron_client.chat.completions.create(
                **kwargs,
                extra_body={"chat_template_kwargs": {"enable_thinking": False}},
            )
        except Exception:
            response = nemotron_client.chat.completions.create(**kwargs)

        if hasattr(response, "usage") and response.usage:
            tokens_used = getattr(response.usage, "total_tokens", 0)

        content = response.choices[0].message.content or ""
        code = self._extract_clean_code(content)
        return code, tokens_used

    def execute(self, code: str, working_dir: Optional[str] = None) -> Tuple[bool, str, str]:
        """
        Execute generated python-docx code in a structured namespace.
        """
        if not code.strip():
            return False, "", "No code provided to execute"

        try:
            ast.parse(code)
        except SyntaxError as syn_err:
            return False, "", f"SyntaxError in generated code: {syn_err}"

        # Extract filename if present in code to provide resilient default doc
        filename_match = re.search(r"filename\s*=\s*['\"]([^'\"]+)['\"]", code)
        doc_inst = None
        if filename_match:
            fname = filename_match.group(1)
            if os.path.exists(fname):
                try:
                    doc_inst = Document(fname)
                except Exception:
                    doc_inst = Document()
            else:
                doc_inst = Document()

        # Setup namespace with all docx and oxml primitives
        namespace: Dict[str, Any] = {
            "os": os,
            "re": re,
            "docx": docx,
            "Document": Document,
            "doc": doc_inst or Document(),
            "Inches": Inches,
            "Pt": Pt,
            "Cm": Cm,
            "Mm": Mm,
            "RGBColor": RGBColor,
            "WD_ALIGN_PARAGRAPH": WD_ALIGN_PARAGRAPH,
            "WD_TABLE_ALIGNMENT": WD_TABLE_ALIGNMENT,
            "parse_xml": parse_xml,
            "OxmlElement": OxmlElement,
            "nsdecls": nsdecls,
            "qn": qn,
            "Path": Path,
        }

        try:
            import openpyxl
            namespace["openpyxl"] = openpyxl
        except ImportError:
            pass

        try:
            import pptx
            namespace["pptx"] = pptx
        except ImportError:
            pass

        stdout_buf = io.StringIO()
        stderr_buf = io.StringIO()

        try:
            with contextlib.redirect_stdout(stdout_buf), contextlib.redirect_stderr(stderr_buf):
                exec(code, namespace)

            captured_out = stdout_buf.getvalue().strip()
            captured_err = stderr_buf.getvalue().strip()
            output_msg = captured_out or "Code executed successfully"
            if captured_err:
                output_msg += f"\nStderr:\n{captured_err}"
            return True, output_msg, ""
        except Exception as e:
            err_details = traceback.format_exc()
            return False, stdout_buf.getvalue().strip(), f"Execution error: {e}\n{err_details}"
