from __future__ import annotations

import os
from typing import Any, Dict, List, Optional
from docx import Document


class DocumentMap:
    """
    Parses and indexes any Word document (.docx) into a clean, unambiguous
    element map with unique paragraph identifiers ([P0], [P1]...) and table identifiers ([T0]...).
    Provides GPS-like coordinates for LLMs to deterministically edit, delete, insert, or restyle.
    """

    def __init__(self, filename: str):
        self.filename = filename
        self.doc: Optional[Document] = None
        self.elements: List[Dict[str, Any]] = []
        self.paragraphs_count = 0
        self.tables_count = 0
        self.load()

    def load(self) -> bool:
        if not os.path.exists(self.filename) or not os.path.isfile(self.filename):
            return False
        try:
            self.doc = Document(self.filename)
            self._parse_elements()
            return True
        except Exception:
            return False

    def _parse_elements(self) -> None:
        self.elements = []
        if not self.doc:
            return

        # Index Paragraphs
        for idx, p in enumerate(self.doc.paragraphs):
            text = p.text.strip()
            style_name = p.style.name if p.style else "Normal"
            self.elements.append({
                "id": f"P{idx}",
                "type": "paragraph",
                "index": idx,
                "style": style_name,
                "text": text,
                "word_count": len(text.split()) if text else 0,
                "is_empty": len(text) == 0,
            })
        self.paragraphs_count = len(self.doc.paragraphs)

        # Index Tables
        for idx, t in enumerate(self.doc.tables):
            rows_count = len(t.rows)
            cols_count = len(t.columns) if t.rows else 0
            headers = []
            if t.rows:
                headers = [c.text.strip() for c in t.rows[0].cells]
            self.elements.append({
                "id": f"T{idx}",
                "type": "table",
                "index": idx,
                "dimensions": f"{rows_count}x{cols_count}",
                "headers": headers,
            })
        self.tables_count = len(self.doc.tables)

    def to_prompt_context(self, max_preview_chars: int = 140) -> str:
        """
        Generate a concise, LLM-ready structural map of the document with paragraph IDs.
        """
        if not self.elements:
            return f"Document '{self.filename}' is empty or could not be parsed."

        lines = [
            f"=== DOCUMENT STRUCTURE MAP: '{self.filename}' ===",
            f"Total Elements: {self.paragraphs_count} Paragraph(s), {self.tables_count} Table(s)",
            "--- Labeled Element Index ---"
        ]

        for el in self.elements:
            if el["type"] == "paragraph":
                pid = el["id"]
                style = el["style"]
                text = el["text"]
                if not text:
                    lines.append(f"[{pid}] (Style: '{style}') [Empty / Blank Paragraph]")
                else:
                    preview = text if len(text) <= max_preview_chars else text[:max_preview_chars] + "..."
                    lines.append(f"[{pid}] (Style: '{style}', {el['word_count']} words) \"{preview}\"")
            elif el["type"] == "table":
                tid = el["id"]
                dims = el["dimensions"]
                headers_str = ", ".join([f"'{h}'" for h in el["headers"][:4]])
                lines.append(f"[{tid}] (Table {dims}) Headers: [{headers_str}]")

        lines.append("=== END OF DOCUMENT MAP ===")
        return "\n".join(lines)


def build_document_map(filename: str) -> DocumentMap:
    """Convenience factory function to inspect and map any .docx file."""
    return DocumentMap(filename)
